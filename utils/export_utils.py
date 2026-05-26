"""Export utilities for TXT, PDF, and DOCX formats."""

import io
from datetime import datetime

from docx import Document
from fpdf import FPDF


def _sanitize_for_pdf(text: str) -> str:
    """Replace characters unsupported by core PDF fonts."""
    replacements = {
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2013": "-",
        "\u2014": "-",
        "\u2022": "*",
        "\u2026": "...",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def export_to_txt(summary: str, title: str = "Summary") -> bytes:
    """Export summary as TXT bytes."""
    header = f"{title}\n{'=' * len(title)}\nGenerated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n\n"
    content = header + summary
    return content.encode("utf-8")


def export_to_pdf(summary: str, title: str = "Summary") -> bytes:
    """Export summary as PDF bytes."""
    
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)

    pdf.add_page()

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 10, _sanitize_for_pdf(title))

    pdf.ln(5)

    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 7, _sanitize_for_pdf(summary))

    pdf_output = pdf.output(dest="S")

    # FIX FOR STREAMLIT
    if isinstance(pdf_output, bytearray):
        pdf_output = bytes(pdf_output)

    elif isinstance(pdf_output, str):
        pdf_output = pdf_output.encode("latin-1")

    return pdf_output


def export_to_docx(summary: str, title: str = "Summary") -> bytes:
    """Export summary as DOCX bytes."""
    document = Document()
    document.add_heading(title, level=1)
    document.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    document.add_paragraph("")
    for paragraph in summary.split("\n"):
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.read()


def build_export_filename(base_name: str, extension: str) -> str:
    """Build a timestamped export filename."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = "".join(c if c.isalnum() or c in ("-", "_") else "_" for c in base_name)
    return f"{safe_name}_{timestamp}.{extension}"
